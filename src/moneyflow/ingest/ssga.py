"""SSGA official daily holdings adapter (M2).

URL pattern (verified 2026-10-01, all 11 tickers return 200):
  https://www.ssga.com/library-content/products/fund-data/etfs/us/
  holdings-daily-us-en-{ticker}.xlsx   (lowercase ticker)

Official daily holdings file, fetched at low frequency with a
descriptive UA. SSGA's Terms of Service have not been independently
reviewed -- no legal conclusion is asserted here. File layout (golden
fixture in tests/fixtures/ssga/): row 2 = file ticker, row 3 =
"As of DD-Mon-YYYY", row 5 = header, rows 6+ = holdings with Weight
in PERCENT.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import openpyxl

from moneyflow.common.http import PoliteClient
from moneyflow.ingest.base import SourceAdapter
from moneyflow.models import Holding

BASE_URL = "https://www.ssga.com/library-content/products/fund-data/etfs/us"
AS_OF_PREFIX = "As of "
AS_OF_FMTS = ("%d-%b-%Y", "%Y-%m-%d", "%d %b %Y")  # primary + fallbacks (m2)
DATA_START_ROW = 6


class SsgaAdapter(SourceAdapter):
    name = "ssga"

    def __init__(self, etf_ticker: str, client: PoliteClient | None = None) -> None:
        self.etf_ticker = etf_ticker.upper()
        # Lazy: parsing never needs HTTP, and building httpx.Client eagerly
        # would parse proxy env vars at import/test time.
        self._client = client
        self._owns_client = client is None

    @property
    def url(self) -> str:
        slug = self.etf_ticker.lower()
        return f"{BASE_URL}/holdings-daily-us-en-{slug}.xlsx"

    def _live_client(self) -> PoliteClient:
        if self._client is None:
            self._client = PoliteClient()
        return self._client

    def fetch_raw(self) -> bytes:
        client = self._live_client()
        try:
            return client.get(self.url).content
        finally:
            if self._owns_client:
                client.close()
                self._client = None

    def parse(self, raw: bytes) -> list[Holding]:
        wb = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
        ws = wb["holdings"]
        file_ticker = ws.cell(row=2, column=2).value
        if file_ticker != self.etf_ticker:
            raise ValueError(
                f"SSGA file ticker mismatch: file says {file_ticker!r}, "
                f"adapter expects {self.etf_ticker!r}"
            )
        as_of_raw = str(ws.cell(row=3, column=2).value or "")
        as_of_clean = as_of_raw.replace(AS_OF_PREFIX, "").strip()
        as_of = None
        for fmt in AS_OF_FMTS:
            try:
                as_of = datetime.strptime(as_of_clean, fmt).date()
                break
            except ValueError:
                continue
        if as_of is None:
            raise ValueError(f"SSGA {self.etf_ticker}: unparseable As-of date {as_of_raw!r}")
        holdings: list[Holding] = []
        seen: set[str] = set()
        for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
            name, ticker, _ident, _sedol, weight = row[0], row[1], row[2], row[3], row[4]
            if not ticker or weight is None:
                continue  # trailing blank rows / disclaimer footnote
            ticker = str(ticker).strip()
            if ticker == "-":
                continue  # cash placeholder rows (money market, USD): not securities
            if ticker in seen:
                raise ValueError(
                    f"SSGA {self.etf_ticker}: duplicate ticker row {ticker!r} "
                    "(glitched file; refusing to double-count)"
                )
            seen.add(ticker)
            holdings.append(
                Holding(
                    as_of=as_of,
                    etf_ticker=self.etf_ticker,
                    ticker=ticker,
                    name=str(name or "").strip(),
                    weight=float(weight) / 100.0,  # file stores percent
                )
            )
        return holdings
